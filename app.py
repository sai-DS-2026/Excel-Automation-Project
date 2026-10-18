import openpyxl as xl
from openpyxl.chart import BarChart, Reference

def process_workflow(filename,sheetname = 'Sheet1'):
    workbook = xl.open(filename)
    mysheet = workbook[sheetname]

    for row in range(2, mysheet.max_row+1):
        cell = mysheet.cell(row,3)
        correct_value = cell.value*0.9
        corrected_value_cell = mysheet.cell(row,4)
        corrected_value_cell.value = correct_value

    values = Reference(mysheet,min_col = 4,max_col = 4,min_row = 2,max_row = mysheet.max_row)

    chart = BarChart()
    chart.add_data(values)
    mysheet.add_chart(chart,'e2')
    workbook.save(filename)
    
